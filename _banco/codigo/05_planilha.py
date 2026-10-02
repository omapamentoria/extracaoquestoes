# Banco MAPA - passo 5: gera catalogo.xlsx (para o Matheus abrir no Excel) a partir do catalogo_debug.json
import os, json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
B = os.path.expanduser("~/mnt/BM/_banco")
L = json.load(open(os.path.join(B, "catalogo_debug.json"), encoding="utf-8"))
wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78")

def aba(ws, cols, linhas, larg=None):
    ws.append(cols)
    for c in ws[1]: c.font = H; c.fill = HF; c.alignment = Alignment(vertical="center", wrap_text=True)
    for l in linhas: ws.append([l.get(c, "") if not isinstance(l, list) else None for c in cols] if isinstance(l, dict) else l)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (larg or {}).get(c, min(40, max(8, len(c) + 2)))

# Resumo
ws = wb.active; ws.title = "Resumo"
ordem = ["ENEM", "SSA", "SIMULADO_ENEM", "SIMULADO_SSA", "FUVEST", "UNICAMP", "UERJ", "UNESP", "DEMAIS",
         "SIMULADO_FUVEST", "SIMULADO_UNICAMP", "SIMULADO_UERJ", "SIMULADO_UNESP", "COMPILACAO", "APOSTILA", "ENCCEJA", "SEM_QUESTOES"]
g = collections.defaultdict(collections.Counter)
for d in L:
    c = g[d["grupo"]]; c["arquivos"] += 1
    if d["status"] == "a extrair":
        c["a extrair"] += 1; c["páginas"] += d["paginas"]; c["questões (estimativa)"] += int(d.get("questoes_estimadas") or 0)
        if not str(d["texto"]).startswith("ok"): c["texto ruim (precisa visão)"] += 1
        if "sem gabarito" in d["obs"]: c["sem gabarito pareado"] += 1
    elif d["status"].startswith("apoio"): c["gabaritos/resoluções"] += 1
    elif "duplicata" in d["status"]: c["duplicatas"] += 1
    elif "sem questões" in d["status"]: c["sem questões"] += 1
cols = ["grupo", "prioridade", "arquivos", "a extrair", "gabaritos/resoluções", "duplicatas", "sem questões", "páginas",
        "questões (estimativa)", "texto ruim (precisa visão)", "sem gabarito pareado"]
prio = {d["grupo"]: d["prioridade"] for d in L}
linhas = [[k, prio.get(k, "")] + [g[k].get(c, 0) for c in cols[2:]] for k in ordem if k in g]
tot = ["TOTAL", ""] + [sum(g[k].get(c, 0) for k in g) for c in cols[2:]]
aba(ws, cols, linhas + [tot], {"grupo": 18})
for c in ws[ws.max_row]: c.font = Font(bold=True)
ws.append([]); ws.append(["Obs.: 'questões (estimativa)' é só uma contagem automática para planejamento; o número real sai na extração."])

C = ["id", "status", "prioridade", "grupo", "tipo_fonte", "vestibular", "instituicao", "ano", "edicao", "fase_etapa", "dia",
     "caderno", "disciplina", "variante", "nivel", "simulado_n", "papel", "gabarito_ids", "resolucao_ids", "padrao_ids",
     "duplicata_de", "tipo_duplicata", "texto", "paginas_problema", "paginas", "questoes_estimadas", "tamanho_mb", "obs",
     "arquivo", "caminho"]
W = {"arquivo": 45, "caminho": 90, "obs": 60, "tipo_duplicata": 30, "texto": 22, "status": 24, "edicao": 18}
aba(wb.create_sheet("A extrair"), C, [d for d in L if d["status"] == "a extrair"], W)
conf = [d for d in L if any(x in d["obs"] for x in ("conferir", "sem gabarito", "não pareado", "cópia da prova")) or
        (d["status"] == "a extrair" and not str(d["texto"]).startswith("ok"))]
aba(wb.create_sheet("Conferir"), C, conf, W)
aba(wb.create_sheet("Catálogo completo"), C, L, W)
wb.save(os.path.join(B, "catalogo.xlsx"))
print("ok:", len(L), "linhas;", len(conf), "para conferir")
